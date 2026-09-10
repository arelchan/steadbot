"""render_check.py — COM 渲染回图(1次) + source|rendered 对比图 + 静态验收.

COM 纪律: ReadOnly + WithWindow=False, 只关自己打开的文件, 绝不 Quit 用户已开实例.
验收: 乱码(??/U+FFFD) / 文本框计数 / 贴图颗粒度(单贴图>35%页面积 WARN, base/ref 豁免).
"""
import json
import ctypes
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SPRITE_AREA_MAX = 0.35


def com_render(pptx_path, out_png, w, h):
    app = pres = None
    preexisting_pids = set()
    pid_probe_ok = False
    instance_pid = None
    isolated = False
    try:
        import pythoncom
        import win32com.client
        try:
            import psutil
            preexisting_pids = {p.pid for p in psutil.process_iter(["name"])
                                if (p.info.get("name") or "").lower() == "powerpnt.exe"}
            pid_probe_ok = True
        except Exception:
            preexisting_pids = set()
        pythoncom.CoInitialize()
        app = win32com.client.DispatchEx("PowerPoint.Application")
        try:
            hwnd = int(app.HWND)
            pid = ctypes.c_ulong()
            ctypes.windll.user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
            instance_pid = int(pid.value)
            isolated = bool(pid_probe_ok and instance_pid and instance_pid not in preexisting_pids)
        except Exception:
            isolated = False
        # WithWindow=False may expose no HWND at all.  In that case the exact
        # before/after process-set delta is still a safe ownership proof.  We
        # only accept a single newly-created POWERPNT PID; ambiguous deltas are
        # deliberately left alone.
        if pid_probe_ok and not isolated:
            try:
                post_pids = {p.pid for p in psutil.process_iter(["name"])
                             if (p.info.get("name") or "").lower() == "powerpnt.exe"}
                created_pids = post_pids - preexisting_pids
                if len(created_pids) == 1:
                    instance_pid = next(iter(created_pids))
                    isolated = True
            except Exception:
                isolated = False
        pres = app.Presentations.Open(str(pptx_path), ReadOnly=True,
                                      Untitled=False, WithWindow=False)
        try:
            pres.Slides(1).Export(str(out_png), "PNG", w, h)
        finally:
            pres.Close()
            pres = None
            # Quit only when the process is proven to be a newly created isolated instance.
            if isolated and app.Presentations.Count == 0:
                app.Quit()
        return {"ok": True, "instance_pid": instance_pid, "isolated": isolated,
                "preexisting_powerpnt_pids": sorted(preexisting_pids)}
    except Exception as ex:  # noqa: BLE001
        print(f"[render] COM unavailable, skip render: {ex}")
        return {"ok": False, "error": str(ex), "instance_pid": instance_pid,
                "isolated": isolated, "preexisting_powerpnt_pids": sorted(preexisting_pids)}
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        try:
            import pythoncom
            pythoncom.CoUninitialize()
        except Exception:
            pass


def validate(run, pptx_path):
    ej = json.loads((run / "elements.json").read_text(encoding="utf-8"))
    prs = Presentation(str(pptx_path))
    s1 = prs.slides[0]
    slide_area = prs.slide_width * prs.slide_height
    issues, texts, n_tb, n_pic, worst = [], [], 0, 0, 0.0
    shape_names = []
    for sh in s1.shapes:
        shape_names.append(sh.name)
        if sh.has_text_frame and sh.text_frame.text.strip():
            n_tb += 1
            texts.append(sh.text_frame.text)
        if sh.shape_type == 13:  # PICTURE
            n_pic += 1
            if sh.name not in ("base", "ref"):
                ratio = (sh.width * sh.height) / slide_area
                worst = max(worst, ratio)
                if ratio > SPRITE_AREA_MAX:
                    issues.append(f"sprite '{sh.name}' covers {ratio:.0%} > {SPRITE_AREA_MAX:.0%} (半页打包?)")
    all_text = "\n".join(texts)
    if "??" in all_text or "�" in all_text:
        issues.append("garbled text detected (??/U+FFFD)")
    expected = [e["text"] for e in ej["elements"] if e["type"] == "text"]
    normalize = lambda s: "".join(str(s).split())  # noqa: E731
    actual_norm = normalize(all_text)
    missing = [t for t in expected if normalize(t) not in actual_norm]
    if missing:
        issues.append(f"missing native text: {missing[:5]}" + (" ..." if len(missing) > 5 else ""))
    n_expect = sum(1 for e in ej["elements"] if e["type"] == "text")
    if n_tb != n_expect:
        issues.append(f"textbox count {n_tb} != elements text count {n_expect}")
    locked = [e["id"] for e in ej["elements"] if e["type"] == "text_in_graphic"]
    demoted = [e["id"] for e in ej["elements"] if e.get("demoted_from") == "card"]
    expected_ids = [e["id"] for e in ej["elements"] if e["type"] in ("text", "card", "graphic")]
    missing_ids = [sid for sid in expected_ids if sid not in shape_names]
    duplicate_ids = sorted({sid for sid in expected_ids if shape_names.count(sid) > 1})
    if missing_ids:
        issues.append(f"missing semantic objects: {missing_ids[:8]}" + (" ..." if len(missing_ids) > 8 else ""))
    if duplicate_ids:
        issues.append(f"duplicate semantic objects: {duplicate_ids[:8]}" + (" ..." if len(duplicate_ids) > 8 else ""))
    n_cards = sum(1 for e in ej["elements"] if e["type"] == "card")
    n_graphics = sum(1 for e in ej["elements"] if e["type"] == "graphic")
    rep = {"textboxes": n_tb, "expected_textboxes": n_expect, "native_text_coverage": 1.0 if not missing else round((len(expected) - len(missing)) / max(1, len(expected)), 3),
           "native_cards": n_cards, "semantic_graphics": n_graphics,
           "semantic_objects": len(expected_ids),
           "semantic_object_coverage": round((len(expected_ids) - len(missing_ids)) / max(1, len(expected_ids)), 3),
           "pictures": n_pic, "max_sprite_area": round(worst, 3),
           "locked_text_regions": locked, "demoted_cards": demoted, "issues": issues}
    (run / "work" / "validate.json").write_text(
        json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return rep


def compare_img(run, rendered):
    src = Image.open(run / "source.png").convert("RGB")
    ren = Image.open(rendered).convert("RGB").resize(src.size)
    bar, w, h = 30, src.width, src.height
    out = Image.new("RGB", (w * 2 + 8, h + bar), "#222222")
    d = ImageDraw.Draw(out)
    d.text((10, 7), "SOURCE", fill="#ffffff")
    d.text((w + 18, 7), "RENDERED", fill="#ffffff")
    out.paste(src, (0, bar))
    out.paste(ren, (w + 8, bar))
    p = run / "work" / "compare.png"
    out.save(p)
    a, b = np.asarray(src, dtype=np.float32), np.asarray(ren, dtype=np.float32)
    mae = float(np.mean(np.abs(a - b)))
    try:
        from skimage.metrics import structural_similarity
        ssim = float(structural_similarity(a.astype(np.uint8), b.astype(np.uint8), channel_axis=2))
    except Exception:
        ssim = None
    return p, {"ssim": round(ssim, 5) if ssim is not None else None, "mae": round(mae, 3)}


def main(run_dir, pptx_path=None):
    run = Path(run_dir)
    (run / "work").mkdir(exist_ok=True)
    ej = json.loads((run / "elements.json").read_text(encoding="utf-8"))
    pptx = Path(pptx_path) if pptx_path else next((p for p in run.glob("*.pptx")), None)
    if not pptx:
        raise SystemExit("[render] no pptx in run dir")
    rendered = run / "work" / "rendered.png"
    com = com_render(pptx.resolve(), rendered.resolve(), ej["img_w"], ej["img_h"])
    rep = validate(run, pptx)
    if com.get("ok"):
        compare_path, metrics = compare_img(run, rendered)
        rep["visual_metrics"] = metrics
        print(f"[render] compare: {compare_path} (SSIM={metrics['ssim']}, MAE={metrics['mae']})")
    rep["com_rendered"] = bool(com.get("ok"))
    rep["com"] = com
    (run / "work" / "validate.json").write_text(
        json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    status = "PASS" if not rep["issues"] else "WARN"
    print(f"[validate] {status}: {rep['textboxes']} textboxes, {rep['pictures']} pictures, "
          f"max sprite {rep['max_sprite_area']:.0%}, locked={rep['locked_text_regions']}")
    for i in rep["issues"]:
        print(f"[validate]   ! {i}")
    return rep


if __name__ == "__main__":
    main(sys.argv[sys.argv.index("--run") + 1] if "--run" in sys.argv else sys.argv[1])
