"""assemble.py — elements.json → 可编辑 pptx (数据驱动一次装配, 无渲染循环).

字号由 PIL 墨水高度反解 (禁启发式); 中文字体三件套 latin+ea+cs; 全 shape 关阴影.
页1 = 底图 + card(原生圆角矩形) + graphic(贴图) + text(原生文本框); 页2 = 原图参考页.
"""
import json
import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SLIDE_W_IN = 13.333
FONT_FILES = {
    "microsoft yahei": (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc"),
    "noto sans sc": (r"C:\Windows\Fonts\NotoSansSC-Regular.ttf", r"C:\Windows\Fonts\NotoSansSC-Bold.ttf"),
}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def font_file(name, bold):
    reg, bd = FONT_FILES.get(name.lower(), FONT_FILES["microsoft yahei"])
    p = Path(bd if bold else reg)
    if not p.exists():
        p = Path(FONT_FILES["microsoft yahei"][1 if bold else 0])
    return str(p)


def ink_box(text, path, px):
    f = ImageFont.truetype(path, px)
    x0, y0, x1, y1 = f.getbbox(text)
    return x0, y0, x1 - x0, y1 - y0  # ink 左偏/顶偏/宽/高


def fit_px(lines, path, box_w, box_h, spacing):
    """反解字号(px): 高度解二分 + 宽度解取小 (血泪#1/#12)."""
    n = len(lines)
    target_h = box_h / (1 + (n - 1) * spacing) if n > 1 else box_h
    # Choose the visually widest line, not the line with most code points.
    ref = max(lines, key=lambda line: ink_box(line, path, 100)[2])
    lo, hi = 4, 400
    while hi - lo > 1:
        mid = (lo + hi) // 2
        if ink_box(ref, path, mid)[3] <= target_h:
            lo = mid
        else:
            hi = mid
    size = lo
    w = ink_box(ref, path, size)[2]
    if w > box_w * 1.03 and w > 0:  # 宽度解: 3% 容差
        size = max(4, int(size * box_w / w))
    return size


def set_typefaces(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)


def no_shadow(shape):
    shape.shadow.inherit = False


def apply_gradient(run, grad):
    """文字渐变填充 (OOXML gradFill)。cutout 自动检测写入 elements.json;
    agent 可手写 {"dir":"v|h","from":"#..","to":"#.."} 或删除该字段禁用。"""
    rPr = run._r.get_or_add_rPr()
    sf = rPr.find(qn("a:solidFill"))
    if sf is not None:
        rPr.remove(sf)
    ang = 5400000 if grad.get("dir") == "v" else 0
    gf = parse_xml(
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<a:gsLst>"
        f'<a:gs pos="0"><a:srgbClr val="{grad["from"].lstrip("#")}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{grad["to"].lstrip("#")}"/></a:gs>'
        "</a:gsLst>"
        f'<a:lin ang="{ang}" scaled="1"/>'
        "</a:gradFill>")
    latin = rPr.find(qn("a:latin"))
    if latin is not None:  # rPr 序列: fill 必须位于 latin/ea/cs 之前
        latin.addprevious(gf)
    else:
        rPr.append(gf)


def set_semantic_metadata(shape, element):
    shape.name = element["id"]
    nodes = shape._element.xpath(".//p:cNvPr")
    if nodes:
        desc = element.get("desc") or element.get("text") or element.get("group") or element["id"]
        nodes[0].set("descr", str(desc))


def main(run_dir, font=None, no_ref=False):
    run = Path(run_dir)
    ej = json.loads((run / "elements.json").read_text(encoding="utf-8"))
    W, H = ej["img_w"], ej["img_h"]
    font = font or ej.get("font", "Microsoft YaHei")
    scale = Emu(int(SLIDE_W_IN * 914400)) / W          # px -> EMU
    px2pt = SLIDE_W_IN * 72.0 / W                      # px -> pt

    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W_IN * 914400))
    prs.slide_height = Emu(int(SLIDE_W_IN * 914400 * H / W))
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def E(v):
        return Emu(int(v * scale))

    if (run / "base.png").exists():
        pic = slide.shapes.add_picture(str(run / "base.png"), 0, 0,
                                       prs.slide_width, prs.slide_height)
        pic.name = "base"

    els = ej["elements"]
    order = {"card": 0, "graphic": 1, "text": 2}
    for e in sorted([e for e in els if e["type"] in order], key=lambda e: order[e["type"]]):
        x0, y0, x1, y1 = e["bbox"]
        if e["type"] == "card":
            sp = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE if e.get("radius") else MSO_SHAPE.RECTANGLE,
                E(x0), E(y0), E(x1 - x0), E(y1 - y0))
            if e.get("radius"):
                sp.adjustments[0] = max(0.0, min(0.5, e["radius"] / max(1, min(x1 - x0, y1 - y0))))
            sp.fill.solid()
            sp.fill.fore_color.rgb = RGBColor.from_string(e.get("fill", "#FFFFFF").lstrip("#"))
            if e.get("stroke"):
                sp.line.fill.solid()
                sp.line.color.rgb = RGBColor.from_string(e["stroke"].lstrip("#"))
                sp.line.width = Pt(float(e.get("stroke_width", 1.0)))
            else:
                sp.line.fill.background()
            no_shadow(sp)
            set_semantic_metadata(sp, e)
        elif e["type"] == "graphic":
            cb = e.get("crop_bbox", [x0, y0, x1, y1])
            pic = slide.shapes.add_picture(
                str(run / e["asset"]), E(cb[0]), E(cb[1]), E(cb[2] - cb[0]), E(cb[3] - cb[1]))
            set_semantic_metadata(pic, e)
        else:  # text
            lines = e["text"].split("\n")
            spacing = e.get("line_spacing", 1.15)
            bold = bool(e.get("bold"))
            ff = font_file(font, bold)
            # OCR 框含标点/引号伸展会比视觉 ink 高 → 字号系统性偏大;
            # agent 框是视觉紧界, cutout 实测的 ink_bbox 更准 (排除 marker/标点),
            # 高度取三者较小值 (宽度仍用 OCR, 含全部字符)。
            ab = e.get("agent_bbox")
            ib = e.get("ink_bbox")
            cands = [y1 - y0] + ([ab[3] - ab[1]] if ab else []) + ([ib[3] - ib[1] + 2] if ib else [])
            eff_h = min(cands)
            # marker 留底图 (marker_stripped) 时, 渲染起点移到字形实际起点,
            # 否则原生文字会盖在保留的 marker 色块上
            if e.get("marker_stripped") and ib:
                x0 = max(x0, ib[0])
            px = fit_px(lines, ff, x1 - x0, eff_h, spacing)
            ix, iy, _, _ = ink_box(lines[0], ff, px)
            # 文本框: 左/顶用 ink 偏移补偿, 宽=bbox 宽(对齐基准), wrap 关
            tb = slide.shapes.add_textbox(E(x0 - ix), E(y0 - iy), E(x1 - x0 + 2 * ix), E((y1 - y0) * 1.4))
            tf = tb.text_frame
            tf.word_wrap = False
            tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                setattr(tf, m, 0)
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = ALIGN.get(e.get("align", "left"), PP_ALIGN.LEFT)
                if len(lines) > 1:
                    p.line_spacing = spacing
                r = p.add_run()
                r.text = line
                r.font.size = Pt(round(px * px2pt * 2) / 2)
                r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(e.get("color", "#1D1D1F").lstrip("#"))
                r.font.name = font
                set_typefaces(r, font)
                grad = e.get("gradient")
                if isinstance(grad, dict):
                    apply_gradient(r, grad)
            no_shadow(tb)
            set_semantic_metadata(tb, e)
            e["fitted_px"] = px

    if not no_ref:
        ref = prs.slides.add_slide(prs.slide_layouts[6])
        pic = ref.shapes.add_picture(str(run / "source.png"), 0, 0,
                                     prs.slide_width, prs.slide_height)
        pic.name = "ref"

    out = run / f"{run.name.split('-', 1)[-1] or run.name}.pptx"
    prs.save(str(out))
    (run / "elements.json").write_text(json.dumps(ej, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"[assemble] saved: {out}")
    return str(out)


if __name__ == "__main__":
    args = sys.argv[1:]
    run = args[args.index("--run") + 1] if "--run" in args else args[0]
    f = args[args.index("--font") + 1] if "--font" in args else None
    main(run, font=f, no_ref="--no-ref" in args)
