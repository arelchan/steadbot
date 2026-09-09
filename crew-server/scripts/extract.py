#!/usr/bin/env python3
"""Read one file as text, whatever it is.

Called by the `see` tool (extensions/see.ts). Prints one JSON object:
  {"kind": "...", "text": "...", "pages": n, "thin": bool, "note": "..."}
`thin` means a PDF whose pages carry almost no text — a scan or an exported deck — so the caller renders the
pages and lets a vision model read them instead.
"""
import json, os, sys, zipfile

MAX = 200_000


def out(kind, text="", **extra):
    if len(text) > MAX:
        text = text[:MAX] + "\n…（已截断）"
    print(json.dumps({"kind": kind, "text": text, **extra}, ensure_ascii=False))
    sys.exit(0)


def pdf(path):
    from pypdf import PdfReader

    r = PdfReader(path)
    parts = []
    for i, page in enumerate(r.pages, 1):
        t = (page.extract_text() or "").strip()
        parts.append(f"--- 第 {i} 页 ---\n{t}" if t else f"--- 第 {i} 页 ---（无文字层）")
    body = "\n\n".join(parts)
    chars = sum(len(p) for p in parts)
    out("pdf", body, pages=len(r.pages), thin=chars / max(1, len(r.pages)) < 120)


def docx(path):
    import docx

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        parts.append(f"[表 {ti}]\n" + "\n".join(rows))
    out("docx", "\n".join(parts))


def pptx(path):
    from pptx import Presentation

    p = Presentation(path)
    parts = []
    for i, slide in enumerate(p.slides, 1):
        lines = [f"--- 第 {i} 页 ---"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            elif shape.shape_type == 13:
                lines.append("[图片]")
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("[备注] " + slide.notes_slide.notes_text_frame.text.strip())
        parts.append("\n".join(lines))
    out("pptx", "\n\n".join(parts), pages=len(p.slides))


def xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= 500:
                rows.append("…（只读了前 500 行）")
                break
            if any(c is not None and str(c).strip() for c in row):
                rows.append(" | ".join("" if c is None else str(c) for c in row))
        parts.append(f"[工作表 {ws.title}]\n" + "\n".join(rows))
    out("xlsx", "\n\n".join(parts), pages=len(wb.worksheets))


def textual(path):
    with open(path, "rb") as f:
        raw = f.read(MAX + 1)
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return out("text", raw.decode(enc))
        except UnicodeDecodeError:
            continue
    out("binary", "", note="不是文本文件")


def main():
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return pdf(path)
        if ext in (".docx", ".doc"):
            return docx(path)
        if ext in (".pptx", ".ppt"):
            return pptx(path)
        if ext in (".xlsx", ".xlsm"):
            return xlsx(path)
        if ext == ".epub" or (zipfile.is_zipfile(path) and ext not in (".zip",)):
            return out("binary", "", note=f"{ext} 需要专门的工具，用 bash 处理")
        return textual(path)
    except Exception as e:  # noqa: BLE001 — the tool turns this into a message for the bot
        out("error", "", note=f"{type(e).__name__}: {e}")


if __name__ == "__main__":
    main()
